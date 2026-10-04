"""Multi-sheet Excel export with light formatting."""
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import config

COLUMNS = ["priority_score", "premium", "title", "organiser", "category", "mode", "eligibility",
           "location", "reg_deadline", "event_date", "prize_pool", "reg_fee",
           "official_url", "discovery_url", "source", "status", "last_verified"]

SOON_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
HEADER_FONT = Font(bold=True)


def _format_sheet(ws):
    for cell in ws[1]:
        cell.font = HEADER_FONT
    for i, col in enumerate(ws.iter_cols(min_row=1), 1):
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 2, 10), 55)
    ws.freeze_panes = "A2"


def export(records: list, export_dir=None) -> Path:
    export_dir = Path(export_dir or config.EXPORT_DIR)
    export_dir.mkdir(parents=True, exist_ok=True)

    df = pd.DataFrame(records)
    for c in COLUMNS:
        if c not in df.columns:
            df[c] = None
    df = df[COLUMNS]

    live = df[df["status"] == "live"]
    premium = live[live["premium"] == 1]

    fname = export_dir / f"competitions_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    with pd.ExcelWriter(fname, engine="openpyxl") as writer:
        premium.to_excel(writer, sheet_name="Premium", index=False)
        live.to_excel(writer, sheet_name="LIVE", index=False)
        df[df["status"] == "broken_link"].to_excel(writer, sheet_name="Broken Links", index=False)
        df[df["status"] == "expired"].to_excel(writer, sheet_name="Expired", index=False)
        for ws in writer.sheets.values():
            _format_sheet(ws)

        ws = writer.sheets["LIVE"]
        dcol = COLUMNS.index("reg_deadline") + 1
        today = datetime.now().date()
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=dcol).value
            if not v:
                continue
            try:
                d = pd.to_datetime(v).date()
            except Exception:
                continue
            if (d - today).days <= 7:
                for c in range(1, len(COLUMNS) + 1):
                    ws.cell(row=row, column=c).fill = SOON_FILL
    return fname
